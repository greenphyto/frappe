# Copyright (c) 2015, Frappe Technologies Pvt. Ltd. and Contributors
# License: MIT. See LICENSE
import datetime
import re
from io import BytesIO

import openpyxl
import xlrd
from openpyxl import load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font
from openpyxl.workbook.child import INVALID_TITLE_REGEX

import frappe
from frappe.utils.html_utils import unescape_html
from frappe.utils import cstr

ILLEGAL_CHARACTERS_RE = re.compile(
	r"[\000-\010]|[\013-\014]|[\016-\037]|\uFEFF|\uFFFE|\uFFFF|[\uD800-\uDFFF]"
)


def get_excel_date_format():
	date_format = frappe.get_system_settings("date_format")
	time_format = frappe.get_system_settings("time_format")

	# Excel-compatible format
	date_format = date_format.replace("mm", "MM")

	return date_format, time_format


# return xlsx file object
def make_xlsx(data, sheet_name, wb=None, column_widths=None, columns=[], bold_list=[]):
	column_widths = column_widths or []
	if wb is None:
		wb = openpyxl.Workbook(write_only=False)

		if wb.active.title == "Sheet":
			wb.remove(wb.active)

	sheet_name_sanitized = INVALID_TITLE_REGEX.sub(" ", sheet_name)
	ws = wb.create_sheet(sheet_name_sanitized, 0)

	for i, column_width in enumerate(column_widths):
		if column_width:
			ws.column_dimensions[get_column_letter(i + 1)].width = column_width

	date_format, time_format = get_excel_date_format()

	for row in data:
		clean_row = []
		for item in row:
			if isinstance(item, str) and (sheet_name not in ["Data Import Template", "Data Export"]):
				value = handle_html(item)
			else:
				value = item

			if isinstance(value, str) and next(ILLEGAL_CHARACTERS_RE.finditer(value), None):
				value = ILLEGAL_CHARACTERS_RE.sub("", value)

			if isinstance(value, datetime.date | datetime.datetime):
				number_format = date_format
				if isinstance(value, datetime.datetime):
					number_format = f"{date_format} {time_format}"

				cell = WriteOnlyCell(ws, value=value)
				cell.number_format = number_format
				clean_row.append(cell)
			else:
				clean_row.append(value)

		ws.append(clean_row)

	# Set bold font for the header (first row)
	for cell in ws[1]:
		cell.font = Font(name="Calibri", bold=True)

	col_right = []
	if columns:
		for i, col in enumerate(columns):
			if col.get("fieldtype") in frappe.model.numeric_fieldtypes:
				col_right.append(i)

	THOUSAND_FORMAT = '#,##0.00;(#,##0.00)'

	for row_index, row in enumerate(ws.iter_rows(values_only=False), start=1):
		for col in col_right:
			if len(row) > col:
				cell = row[col]
				cell.alignment = Alignment(horizontal='right')
				if row_index > 1:
					cell.number_format = THOUSAND_FORMAT

		if row_index in bold_list:
			for cell in row:
				cell.font = Font(bold=True)

	xlsx_file = BytesIO()
	wb.save(xlsx_file)
	return xlsx_file

def handle_html(data):
	from frappe.core.utils import html2text

	# return if no html tags found
	data = frappe.as_unicode(data)

	if "<" not in data or ">" not in data:
		return data

	h = unescape_html(data or "")

	try:
		value = html2text(h, strip_links=True, wrap=False)
	except Exception:
		# unable to parse html, send it raw
		return data

	value = ", ".join(value.split("  \n"))
	value = " ".join(value.split("\n"))
	value = ", ".join(value.split("# "))

	return value

def read_xlsx_file_from_attached_file(file_url=None, fcontent=None, filepath=None):
	if file_url:
		_file = frappe.get_doc("File", {"file_url": file_url})
		filename = _file.get_full_path()
	elif fcontent:
		filename = BytesIO(fcontent)
	elif filepath:
		filename = filepath
	else:
		return

	rows = []
	wb1 = load_workbook(filename=filename, read_only=True, data_only=True)
	ws1 = wb1.active
	for row in ws1.iter_rows():
		tmp_list = []
		for cell in row:
			tmp_list.append(cell.value)
		rows.append(tmp_list)
	return rows

def read_xls_file_from_attached_file(content):
	book = xlrd.open_workbook(file_contents=content)
	sheets = book.sheets()
	sheet = sheets[0]
	rows = []
	for i in range(sheet.nrows):
		rows.append(sheet.row_values(i))
	return rows

def build_xlsx_response(data, filename):
	xlsx_file = make_xlsx(data, filename)
	# write out response as a xlsx type
	frappe.response["filename"] = filename + ".xlsx"
	frappe.response["filecontent"] = xlsx_file.getvalue()
	frappe.response["type"] = "binary"
